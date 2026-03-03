from importlib.resources import path
import os
from pypdf import PdfReader
from pdf2image import convert_from_path


def extract_pages_from_pdf(file_path: str) -> list[dict]:
    if(not os.path.exists(file_path)):
        raise FileNotFoundError("file not found")
    
    reader = PdfReader(file_path)
    records = []
    for pageno,page in enumerate(reader.pages,start = 1):
        text = (page.extract_text() or "").strip()
        
        records.append({
                "text": text,
                "doc_name": os.path.basename(file_path),
                "page": pageno,
                "source": "pdf"
            })
            
    return records




from pptx import Presentation

def extract_slides_from_ppt(file_path: str):
    prs = Presentation(file_path)
    records = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    slide_text.append(text)

        full_text = "\n".join(slide_text).strip()

        if full_text:
            records.append({
                "text": full_text,
                "doc_name": os.path.basename(file_path),
                "page": slide_num,     # slide number
                "source": "ppt"
            })

    return records


def load_documents(file_paths):
    all_records = []

    for path in file_paths:
        ext = path.lower()

        if ext.endswith(".pdf"):
            all_records.extend(extract_pages_from_pdf(path))

        elif ext.endswith(".pptx"):
            all_records.extend(extract_slides_from_ppt(path))

        elif ext.endswith(".ppt"):
            raise ValueError(
                f"Old .ppt files are not supported. Please convert '{os.path.basename(path)}' to .pptx."
            )

        else:
            raise ValueError(f"Unsupported file type: {path}")

    return [r for r in all_records if r["text"].strip()]


        