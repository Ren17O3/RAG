import os
from pypdf import PdfReader
from src.schemas.models import DocumentRecord
import re
import uuid
from pptx import Presentation


def clean_text(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()


def extract_pages_from_pdf(file_path: str, original_name: str) -> list[DocumentRecord]:
    if not os.path.exists(file_path):
        raise FileNotFoundError("file not found")

    reader = PdfReader(file_path)
    records = []
    document_id = str(uuid.uuid4())
    for pageno, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")

        records.append(
            DocumentRecord(
                text=text,
                doc_name=original_name,
                page=pageno,
                source="pdf",
                document_id=document_id,
                metadata={"file_type": "pdf", "file_name": original_name},
            )
        )

    return records


def extract_slides_from_ppt(file_path: str, original_name: str) -> list[DocumentRecord]:
    prs = Presentation(file_path)
    records = []
    document_id = str(uuid.uuid4())
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text:
                    slide_text.append(text)

        full_text = clean_text("\n".join(slide_text))

        if full_text:
            records.append(
                DocumentRecord(
                    text=full_text,
                    doc_name=original_name,
                    page=slide_num,
                    source="ppt",
                    document_id=document_id,
                    metadata={
                        "file_type": "ppt",
                        "file_name": original_name,
                    },
                )
            )

    return records


def load_documents(files: list[dict]) -> list[DocumentRecord]:
    all_records = []

    for file_info in files:
        path = file_info["temp_path"]
        original_name = file_info["original_name"]
        ext = path.lower()

        if ext.endswith(".pdf"):
            all_records.extend(extract_pages_from_pdf(path, original_name))

        elif ext.endswith(".pptx"):
            all_records.extend(extract_slides_from_ppt(path, original_name))

        elif ext.endswith(".ppt"):
            raise ValueError(
                f"Old .ppt files are not supported. Please convert '{os.path.basename(path)}' to .pptx."
            )

        else:
            raise ValueError(f"Unsupported file type: {path}")

    return [r for r in all_records if r.text.strip()]
